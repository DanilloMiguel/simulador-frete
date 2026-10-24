import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

import frete
import gerar_dados


# Formato original da tabela (só região, sem faixa de peso e sem %): continua aceito.
COLUNAS_FORMATO_ANTIGO = [
    "transportadora", "origem", "regiao", "destino_tipo", "custo_base", "custo_por_kg", "prazo_dias",
]


def tabela(*linhas, regiao="Sudeste"):
    """(transportadora, origem, destino_tipo, custo_base, custo_por_kg, prazo_dias).

    A região é a mesma para todas as linhas (por padrão, Sudeste, a de SP e MG).
    """
    dados = [(t, o, regiao, d, b, k, p) for (t, o, d, b, k, p) in linhas]
    return pd.DataFrame(dados, columns=COLUNAS_FORMATO_ANTIGO)


def pedidos(*linhas):
    """(pedido_id, origem, destino_tipo, uf, peso_kg, valor_mercadoria)"""
    return pd.DataFrame(linhas, columns=frete.COLUNAS_PEDIDOS)


P1 = ("P1", "CD", "Capital", "SP", 2.0, 100.0)


# ------------------------------------------------------------------ simulação
def test_escolhe_a_mais_barata():
    t = tabela(("A", "CD", "Capital", 10, 2, 5), ("B", "CD", "Capital", 9, 2, 5))
    r = frete.simular(pedidos(P1), t)
    assert r.loc[0, "transportadora"] == "B"
    assert r.loc[0, "custo_frete"] == 13.0


def test_empate_de_custo_decide_pelo_menor_prazo():
    t = tabela(("A", "CD", "Capital", 10, 2, 5), ("B", "CD", "Capital", 8, 3, 3))
    r = frete.simular(pedidos(P1), t)  # ambas custam 14
    assert r.loc[0, "transportadora"] == "B"


def test_respeita_prazo_maximo():
    t = tabela(("A", "CD", "Capital", 5, 1, 8), ("B", "CD", "Capital", 20, 2, 3))
    assert frete.simular(pedidos(P1), t, prazo_max=5).loc[0, "transportadora"] == "B"
    sem = frete.simular(pedidos(P1), t, prazo_max=2)
    assert sem.loc[0, "transportadora"] == frete.SEM_COBERTURA
    assert pd.isna(sem.loc[0, "custo_frete"])


def test_origem_fora_da_tabela_fica_sem_cobertura():
    t = tabela(("A", "OUTRO", "Capital", 10, 2, 5))
    r = frete.simular(pedidos(P1), t)
    assert r.loc[0, "transportadora"] == frete.SEM_COBERTURA


def test_regiao_vem_da_uf_do_pedido():
    sul = tabela(("A", "CD", "Capital", 10, 2, 5), regiao="Sul")
    assert frete.simular(pedidos(("P1", "CD", "Capital", "pr", 2.0, 100.0)), sul).loc[0, "transportadora"] == "A"
    assert frete.simular(pedidos(P1), sul).loc[0, "transportadora"] == frete.SEM_COBERTURA  # SP é Sudeste


def test_todas_as_ufs_tem_regiao():
    assert len(frete.REGIAO_POR_UF) == 27
    assert set(frete.REGIAO_POR_UF.values()) == set(frete.ORDEM_REGIOES)


def test_media_ponderada_nao_e_media_das_medias():
    t = tabela(("A", "CD", "Capital", 10, 0, 2), ("A", "CD", "Interior", 20, 0, 4))
    ped = pedidos(
        ("P1", "CD", "Capital", "SP", 1.0, 50.0),
        ("P2", "CD", "Interior", "MG", 1.0, 50.0),
        ("P3", "CD", "Interior", "MG", 1.0, 50.0),
        ("P4", "CD", "Interior", "MG", 1.0, 50.0),
    )
    res = frete.resumo(frete.simular(ped, t))
    total = res[res["origem"] == "TOTAL"].iloc[0]
    assert total["custo_medio"] == pytest.approx(17.5)  # 70 / 4
    assert total["custo_medio"] != pytest.approx(15.0)  # média simples de 10 e 20


# ------------------------------------------------------------------- validação
def test_pedido_id_repetido_gera_erro():
    t = tabela(("A", "CD", "Capital", 10, 2, 5))
    with pytest.raises(ValueError, match="repetido"):
        frete.simular(pedidos(P1, P1), t)


def test_coluna_faltando_gera_erro():
    t = tabela(("A", "CD", "Capital", 10, 2, 5)).drop(columns="prazo_dias")
    with pytest.raises(ValueError, match="prazo_dias"):
        frete.simular(pedidos(P1), t)


def test_valor_invalido_gera_erro():
    t = tabela(("A", "CD", "Capital", "dez", 2, 5))
    with pytest.raises(ValueError, match="inválido"):
        frete.simular(pedidos(P1), t)


def test_uf_desconhecida_gera_erro():
    t = tabela(("A", "CD", "Capital", 10, 2, 5))
    with pytest.raises(ValueError, match="XX"):
        frete.simular(pedidos(("P1", "CD", "Capital", "XX", 2.0, 100.0)), t)


# --------------------------------------------------------------------- ranking
def test_ranking_top3_e_variacao_contra_o_seguinte():
    t = tabela(
        ("A", "CD", "Capital", 10, 0, 3),
        ("B", "CD", "Capital", 12, 0, 4),
        ("C", "CD", "Capital", 15, 0, 5),
        ("D", "CD", "Capital", 20, 0, 6),
    )
    ped = pedidos(P1, ("P2", "CD", "Capital", "MG", 3.0, 100.0))
    r = frete.ranking(frete.calcular_opcoes(ped, t), ["regiao"])
    assert list(r["transportadora"]) == ["A", "B", "C"]  # D fica fora do Top 3
    assert list(r["posicao"]) == [1, 2, 3]
    assert r.loc[0, "var_pct"] == pytest.approx((10 / 12 - 1) * 100)
    assert r.loc[1, "var_pct"] == pytest.approx((12 / 15 - 1) * 100)
    assert pd.isna(r.loc[2, "var_pct"])


def test_ranking_usa_media_ponderada_dos_pedidos():
    t = tabela(("A", "CD", "Capital", 10, 0, 2), ("A", "CD", "Interior", 20, 0, 4))
    ped = pedidos(
        ("P1", "CD", "Capital", "SP", 1.0, 50.0),
        ("P2", "CD", "Interior", "SP", 1.0, 50.0),
        ("P3", "CD", "Interior", "SP", 1.0, 50.0),
        ("P4", "CD", "Interior", "SP", 1.0, 50.0),
    )
    r = frete.ranking(frete.calcular_opcoes(ped, t), ["regiao"])
    assert r.loc[0, "frete"] == pytest.approx(17.5)  # 70 / 4, e não (10 + 20) / 2
    assert r.loc[0, "prazo"] == pytest.approx(3.5)  # (2 + 4 + 4 + 4) / 4


def test_ranking_empate_decide_pelo_prazo():
    t = tabela(("A", "CD", "Capital", 10, 0, 5), ("B", "CD", "Capital", 10, 0, 3))
    r = frete.ranking(frete.calcular_opcoes(pedidos(P1), t), ["uf"])
    assert list(r["transportadora"]) == ["B", "A"]


def test_ranking_geral_sem_agrupamento():
    t = tabela(("A", "CD", "Capital", 10, 0, 5), ("B", "CD", "Capital", 12, 0, 3))
    r = frete.ranking(frete.calcular_opcoes(pedidos(P1), t), [])
    assert list(r["transportadora"]) == ["A", "B"]


def test_dados_de_exemplo_sao_coerentes():
    ped = gerar_dados.gerar_pedidos(np.random.default_rng(1))
    tab = gerar_dados.gerar_tabela()
    assert set(gerar_dados.PESO_UF) == set(frete.REGIAO_POR_UF)
    faixas = len(gerar_dados.FAIXAS) + 1  # + linha de excedente
    por_regiao = 4 * 3 * 5 * 2 * faixas  # transportadoras × origens × regiões × destinos × faixas
    por_uf = len(gerar_dados.EXCECOES_UF) * 3 * 2 * faixas
    assert len(tab) == por_regiao + por_uf
    assert tab["peso_max"].dropna().max() == 30  # faixas até 30 kg + excedente
    assert (frete.simular(ped, tab)["transportadora"] != frete.SEM_COBERTURA).all()


# ----------------------------------------------------------------------- excel
def test_excel_tem_as_quatro_abas():
    t = tabela(("A", "CD", "Capital", 10, 2, 5))
    r = frete.simular(pedidos(P1), t)
    dados = frete.exportar_excel(r, frete.resumo(r), frete.por_transportadora(r), t)
    wb = openpyxl.load_workbook(io.BytesIO(dados))
    assert wb.sheetnames == ["Resultado", "Resumo", "Transportadoras", "Parametros"]


def test_excel_usa_rotulos_e_mantem_parametros_reutilizaveis():
    t = tabela(("A", "CD", "Capital", 10, 2, 5))
    r = frete.simular(pedidos(P1), t)
    dados = frete.exportar_excel(r, frete.resumo(r), frete.por_transportadora(r), t)
    wb = openpyxl.load_workbook(io.BytesIO(dados))
    cabecalho = lambda aba: [c.value for c in wb[aba][1]]
    assert "Frete (R$)" in cabecalho("Resultado")
    assert "Custo médio (R$)" in cabecalho("Resumo")
    assert cabecalho("Parametros") == frete.COLUNAS_TABELA  # pode voltar como entrada


# ------------------------------------------------------- valor de mercadoria
def test_aceita_o_nome_antigo_valor_pedido():
    antigo = pedidos(P1).rename(columns={"valor_mercadoria": "valor_pedido"})
    t = tabela(("A", "CD", "Capital", 10, 2, 5))
    assert "valor_mercadoria" in frete.normalizar_pedidos(antigo).columns
    assert frete.simular(antigo, t).loc[0, "valor_mercadoria"] == 100.0


def test_rotulo_e_planilha_modelo_falam_em_valor_de_mercadoria():
    assert frete.ROTULOS["valor_mercadoria"] == "Valor de mercadoria (R$)"
    wb = openpyxl.load_workbook(io.BytesIO(frete.gerar_template_excel()))
    assert "valor_mercadoria" in [c.value for c in wb["Pedidos"][1]]
    assert "valor_pedido" not in [c.value for c in wb["Pedidos"][1]]


# ------------------------------------------------------------------ mapa por UF
def test_por_uf_frete_medio_e_ponderado():
    t = tabela(("A", "CD", "Capital", 10, 0, 2), ("A", "CD", "Interior", 20, 0, 4))
    ped = pedidos(
        ("P1", "CD", "Capital", "SP", 1.0, 50.0),
        ("P2", "CD", "Interior", "SP", 1.0, 50.0),
        ("P3", "CD", "Interior", "SP", 1.0, 50.0),
        ("P4", "CD", "Capital", "MG", 1.0, 50.0),
    )
    uf = frete.por_uf(frete.simular(ped, t)).set_index("uf")
    assert uf.loc["SP", "pedidos"] == 3
    assert uf.loc["SP", "frete_medio"] == pytest.approx(50 / 3)  # 50 / 3 pedidos
    assert uf.loc["MG", "frete_medio"] == pytest.approx(10)
    assert uf["pct"].sum() == pytest.approx(100)


def test_por_uf_ignora_pedidos_sem_cobertura():
    t = tabela(("A", "OUTRO", "Capital", 10, 0, 2))
    assert frete.por_uf(frete.simular(pedidos(P1), t)).empty


# -------------------------------------------------------- % custo de frete
def test_pct_custo_frete_e_frete_sobre_valor_de_mercadoria():
    t = tabela(("A", "CD", "Capital", 10, 0, 2))
    ped = pedidos(("P1", "CD", "Capital", "SP", 1.0, 100.0), ("P2", "CD", "Capital", "SP", 1.0, 25.0))
    pct = frete.pct_custo_frete(frete.simular(ped, t)).set_index("pedido_id")["pct_frete"]
    assert pct["P1"] == pytest.approx(10.0)  # 10 / 100
    assert pct["P2"] == pytest.approx(40.0)  # 10 / 25


def test_indicadores_media_simples_difere_da_ponderada():
    t = tabela(("A", "CD", "Capital", 10, 0, 2))
    ped = pedidos(
        ("P1", "CD", "Capital", "SP", 1.0, 100.0),  # 10%
        ("P2", "CD", "Capital", "SP", 1.0, 25.0),  # 40%
        ("P3", "CD", "Capital", "SP", 1.0, 50.0),  # 20%
    )
    ind = frete.indicadores_pct_frete(frete.pct_custo_frete(frete.simular(ped, t)))
    assert ind["media"] == pytest.approx((10 + 40 + 20) / 3)
    assert ind["mediana"] == pytest.approx(20.0)
    assert ind["maximo"] == pytest.approx(40.0)
    assert ind["ponderado"] == pytest.approx(30 / 175 * 100)  # soma do frete / soma do valor


def test_indicadores_sem_pedidos_nao_quebram():
    t = tabela(("A", "OUTRO", "Capital", 10, 0, 2))
    ind = frete.indicadores_pct_frete(frete.pct_custo_frete(frete.simular(pedidos(P1), t)))
    assert all(pd.isna(v) for v in ind.values())


# ------------------------------------------- UF, faixas de peso e GRIS + ADV
def tabela_nova(*linhas):
    """(transportadora, origem, uf, regiao, destino_tipo, peso_min, peso_max,
    custo_base, custo_por_kg, gris_adv_pct, prazo_dias). None = célula vazia."""
    return pd.DataFrame(linhas, columns=frete.COLUNAS_TABELA)


FAIXAS_A = [  # transportadora A, CD → SP capital, no formato da GEOCOM
    ("A", "CD", "SP", None, "Capital", 0.000, 0.250, 10.0, 0, 0, 3),
    ("A", "CD", "SP", None, "Capital", 0.251, 0.500, 12.0, 0, 0, 3),
    ("A", "CD", "SP", None, "Capital", 0.501, 30.000, 50.0, 0, 0, 3),
]


def frete_de(peso, *linhas, valor=100.0):
    r = frete.simular(pedidos(("P1", "CD", "Capital", "SP", peso, valor)), tabela_nova(*linhas))
    return r.loc[0, "custo_frete"], r.loc[0, "transportadora"]


def test_faixa_de_peso_escolhe_a_faixa_certa():
    assert frete_de(0.250, *FAIXAS_A)[0] == 10.0
    assert frete_de(0.300, *FAIXAS_A)[0] == 12.0
    assert frete_de(0.000, *FAIXAS_A)[0] == 10.0


def test_peso_entre_duas_faixas_vai_para_a_de_cima():
    assert frete_de(0.2505, *FAIXAS_A)[0] == 12.0  # entre 0,250 e 0,251


def test_acima_da_ultima_faixa_sem_excedente_fica_sem_cobertura():
    assert frete_de(31.0, *FAIXAS_A)[1] == frete.SEM_COBERTURA


def test_linha_de_excedente_cobra_por_kg_acima_de_30():
    excedente = ("A", "CD", "SP", None, "Capital", 30.0, None, 50.0, 2.0, 0, 3)
    assert frete_de(32.5, *FAIXAS_A, excedente)[0] == 55.0  # 50 + 2 × 2,5
    assert frete_de(30.0, *FAIXAS_A, excedente)[0] == 50.0  # ainda na faixa até 30


def test_gris_adv_e_percentual_do_valor_de_mercadoria():
    linha = ("A", "CD", "SP", None, "Capital", 0, None, 10.0, 0, 0.5, 3)  # 0,5%
    custo, _ = frete_de(2.0, linha, valor=400.0)
    assert custo == 12.0  # 10 + 0,5% × 400
    r = frete.simular(pedidos(("P1", "CD", "Capital", "SP", 2.0, 400.0)), tabela_nova(linha))
    assert r.loc[0, "gris_adv_valor"] == 2.0


def test_linha_da_uf_tem_prioridade_sobre_a_da_regiao():
    t = tabela_nova(
        ("A", "CD", None, "Sudeste", "Capital", 0, None, 20.0, 0, 0, 3),
        ("A", "CD", "SP", None, "Capital", 0, None, 15.0, 0, 0, 3),
    )
    ped = pedidos(("P1", "CD", "Capital", "SP", 1.0, 100.0), ("P2", "CD", "Capital", "MG", 1.0, 100.0))
    r = frete.simular(ped, t).set_index("pedido_id")
    assert r.loc["P1", "custo_frete"] == 15.0  # SP usa a linha da UF
    assert r.loc["P2", "custo_frete"] == 20.0  # MG continua com a da região


def test_linha_so_com_regiao_vale_para_todas_as_ufs_da_regiao():
    t = tabela_nova(("A", "CD", None, "Norte", "Capital", 0, None, 30.0, 0, 0, 9))
    prep = frete.preparar_tabela(t)
    assert sorted(prep["uf"]) == sorted(u for u, r in frete.REGIAO_POR_UF.items() if r == "Norte")


def test_formato_novo_e_antigo_dao_o_mesmo_frete():
    antigo = tabela(("A", "CD", "Capital", 10, 2, 5))
    novo = tabela_nova(("A", "CD", None, "Sudeste", "Capital", None, None, 10, 2, None, 5))
    assert frete.simular(pedidos(P1), antigo).equals(frete.simular(pedidos(P1), novo))


def test_sensibilidade_ajusta_tambem_o_gris_adv():
    linha = ("A", "CD", "SP", None, "Capital", 0, None, 10.0, 0, 1.0, 3)  # 10 + 1% × 100 = 11
    r = frete.sensibilidade(pedidos(P1), tabela_nova(linha), "A", 10)
    assert r.loc[0, "custo_frete"] == 12.1  # tudo +10%


@pytest.mark.parametrize(
    "linha, trecho",
    [
        (("A", "CD", None, None, "Capital", 0, None, 10, 0, 0, 3), "preencha uf ou regiao"),
        (("A", "CD", "SP", "Sul", "Capital", 0, None, 10, 0, 0, 3), "não combinam"),
        (("A", "CD", "XX", None, "Capital", 0, None, 10, 0, 0, 3), "UF não reconhecida"),
        (("A", "CD", "SP", None, "Capital", 0, None, 10, 0, "0,3%", 3), "gris_adv_pct"),
        (("A", "CD", "SP", None, "Capital", 5, 1, 10, 0, 0, 3), "peso_min maior"),
        (("A", "CD", "SP", None, "Capital", 0, None, None, 0, 0, 3), "Se não houver custo, use 0"),
    ],
)
def test_erros_de_preenchimento_da_tabela(linha, trecho):
    with pytest.raises(ValueError, match=trecho):
        frete.simular(pedidos(P1), tabela_nova(linha))


def test_faixa_repetida_gera_erro():
    t = tabela_nova(*FAIXAS_A, FAIXAS_A[0])
    with pytest.raises(ValueError, match="repetidas"):
        frete.simular(pedidos(P1), t)


def test_tabela_sem_uf_nem_regiao_gera_erro():
    t = tabela(("A", "CD", "Capital", 10, 2, 5)).drop(columns="regiao")
    with pytest.raises(ValueError, match="uf, a coluna regiao"):
        frete.simular(pedidos(P1), t)


def test_regiao_aceita_maiusculas_e_minusculas():
    t = tabela_nova(("A", "CD", None, "sudeste", "Capital", 0, None, 10, 0, 0, 3))
    assert frete_de(1.0, *t.itertuples(index=False))[1] == "A"


def test_planilha_modelo_roda_no_simulador():
    xls = pd.read_excel(io.BytesIO(frete.gerar_template_excel()), sheet_name=None)
    r = frete.simular(xls["Pedidos"], xls["Tabela de frete"])
    assert len(r) == 2
